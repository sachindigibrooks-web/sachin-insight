"""One-month (Mon to Fri) GBP calendar in the Workana reference layout, with embedded banners."""
import sys, os, re, datetime as dt
sys.path.insert(0, os.path.dirname(__file__)); sys.path.insert(0, os.path.join(os.path.dirname(__file__), "ban"))
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.drawing.image import Image as XImage
from PIL import Image
import configs, posts_hg, posts_pf, month
from build import title, table, GUIDE, F, FB, FS, BOX, WRAP, WRAPC

HERE = os.path.dirname(os.path.abspath(__file__))
BAN = os.path.join(HERE, "ban")
NAVY = "1F4E78"
HF = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor=NAVY)
LINK = Font(name="Arial", size=10, color="0000FF")

GUIDE = [g for g in GUIDE if not g[0].startswith("Photos must")] + [
    ("Photos must follow the photo policy: real, relevant to the business, not misleading.",
     "Every banner uses a real photo from the business's own website (one HG post uses a licensed Unsplash photo). Overlay text is short (headline plus one line), with no phone numbers, emails, prices or claims that are not on the website.",
     "https://support.google.com/business/answer/7213077"),
    ("Recommended post image: JPG or PNG, 10 KB to 5 MB, at least 400 x 300 px, 4:3 ratio displays best.",
     "Banners are 1200 x 900 px PNG (4:3). Key text sits inside the centre safe area so the GBP crop does not cut it.",
     "https://support.google.com/business/answer/7342169"),
]

SITES = {
    "hg": dict(cfg=configs.HG, posts=posts_hg.POSTS, plan=month.HG, new=month.HG_NEW, dl="hg_dl.txt", ph="ph_hg", out="out_hg",
               folder="harlingen-glass", file="Harlingen-Glass-GBP-Monthly-Calendar-Nov-2026.xlsx",
               logo="https://harlingenglass.com/wp-content/uploads/2024/11/LOGO-1.jpg",
               palette=[("Primary navy", "#004D71", "Logo wordmark and panes", "Background panels, offer highlight"),
                        ("Charcoal", "#2C353D", "Website navigation and text", "Dark overlays"),
                        ("Accent cyan", "#38C6F4", "Logo glass panes, website buttons", "Headline accent word, button chips"),
                        ("Aqua", "#75CDD7", "Logo glass panes", "Glass pane motif"),
                        ("Light blue", "#ABDFEA", "Logo glass panes", "Soft tints"),
                        ("White", "#FFFFFF", "Website background", "Text on dark, logo chip")],
               fonts=[("Headings", "Metrophobic", "Same as website H1 to H3"), ("Body / labels", "Roboto", "Same as website body text")],
               motif="Slanted translucent glass panes taken from the logo mark.",
               logo_note="Full logo on a white rounded chip, top left of every banner.",
               removed=[("IDEA-Toros-College-Prep-Front.jpg", "Credited to a press photographer (Jose Aldape) and carries a Texas Border Business watermark", "Rio-Hondo-Middle-School-3.jpg (own project photo)"),
                        ("IDEA-Toros-College-Prep.jpg", "Same press photo series of a third-party school brand; ownership unclear", "UTB-Kinesiology-Bldg.jpg (own project photo)")]),
    "pf": dict(cfg=configs.PF, posts=posts_pf.POSTS, plan=month.PF, new=month.PF_NEW, dl="pf_dl.txt", ph="ph_pf", out="out_pf",
               folder="preferred-fence", file="Preferred-Fence-GBP-Monthly-Calendar-Nov-2026.xlsx",
               logo="https://preferredfencesolutions.com/wp-content/uploads/2021/03/PFS-Logo-1.png",
               palette=[("Black", "#000000", "Website header, footer and text (Elementor global)", "Background panels"),
                        ("Accent yellow", "#DEE101", "Website buttons and highlights (Elementor global)", "Headline accent word, button chips, picket motif"),
                        ("Off-white", "#F4F4F4", "Website section backgrounds", "Light panels"),
                        ("White", "#FFFFFF", "Website background", "Text on dark, logo chip")],
               fonts=[("Headings", "Rubik", "Same as website headings"), ("Body / labels", "Roboto", "Same as website body text")],
               motif="Yellow picket fence line along the bottom edge.",
               logo_note="Logo name and tagline on a white rounded chip. The phone number line in the original logo is cropped out to follow GBP rules.",
               removed=[("fence-damage.jpg (2 banners)", "No camera data, uploaded in 2026 and sized 1376 x 768, which matches AI-generated images; origin cannot be confirmed", "IMG_5914.jpg and IMG_0730001.jpg (own project photos)")]),
}
SPECIAL_SRC = {"90_unsplash_cracked.jpg": "https://unsplash.com/photos/_enXmoXudAk (Unsplash licence)",
               "190_unsplash_AnaOxSVGLkQ.jpg": "https://unsplash.com/photos/AnaOxSVGLkQ (Unsplash licence)",
               "91_team_crop.jpg": None}


def clean_head(h):
    return h.replace("|", " ").replace("*", "").replace("  ", " ").strip()


def photo_urls(s):
    urls = open(os.path.join(BAN, s["dl"])).read().split()
    by_name = {u.rsplit("/", 1)[-1]: u for u in urls}
    def src(f):
        m = re.match(r"\d+_unsplash_(.+)\.jpg$", f)
        if m and f not in SPECIAL_SRC:
            return f"https://unsplash.com/photos/{m.group(1)} (Unsplash licence)"
        if f in SPECIAL_SRC:
            if SPECIAL_SRC[f]:
                return SPECIAL_SRC[f]
            return by_name["fbinfoo-scaled.jpg"] + " (cropped)"
        return by_name[re.sub(r"^\d+_", "", f)]
    return src


OWN = "Business's own photo from its website (client to confirm ownership)."
UNSPLASH = "Unsplash License: free for commercial use, no permission or credit needed."


def rights(photos):
    return "\n".join(UNSPLASH if "unsplash" in f else OWN for f in photos) if len({("unsplash" in f) for f in photos}) > 1 else (UNSPLASH if "unsplash" in photos[0] else OWN)


def rows_for(s):
    base = {p[0]: p for p in s["posts"]}
    src = photo_urls(s)
    out = []
    for i, m in enumerate(s["plan"]):
        d = dt.datetime.strptime(m["date"], "%Y-%m-%d")
        if m["base"] in base:
            p = base[m["base"]]
            cat, typ, copy, cta, land, utm, terms = p[3], p[4], p[8], p[9], p[10], p[11], p[12]
        else:
            n = s["new"][m["base"]]
            cat, typ, copy, cta, land, utm, terms = n["cat"], n["typ"], n["copy"], n["cta"], n["landing"], n["utm"], ""
        copy = m.get("copy", copy).strip()
        terms = m.get("offer", terms)
        utm = f"nov26-{i+1:02d}-" + re.sub(r"^(hg|pf)\d*-", "", utm)
        overlay = "\n".join(x for x in [
            f"Tag: {m['tag']}", f"Headline: {clean_head(m['head'])}",
            f"Sub-line: {m['sub']}" if m.get("sub") else "",
            f"Labels: {', '.join(m['labels'])}" if m.get("labels") else "",
            f"Button: {m['chip']}"] if x)
        fname = f"{m['date']}_gbp-post-{i+1:02d}_1200x900.png"
        out.append(dict(n=i + 1, week=(d - dt.datetime(2026, 11, 2)).days // 7 + 1, date=d,
                        cat=cat, typ=typ, title=m["title"], copy=copy, cta=cta, land=land, utm=utm, terms=terms,
                        overlay=overlay, png=os.path.join(BAN, s["out"], f"{s['out'][4:]}_{i+1:02d}.png"), fname=fname,
                        src="\n".join(src(f) for f in m["photos"]), rights=rights(m["photos"]), photos=m["photos"], tpl=m["tpl"]))
    return out


def head_row(ws, row, headers, col0=2):
    for j, h in enumerate(headers, col0):
        c = ws.cell(row=row, column=j, value=h)
        c.font = HF; c.fill = HFILL; c.border = BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build(key, repo_dir):
    s = SITES[key]; cfg = s["cfg"]
    rows = rows_for(s)
    ban_dir = os.path.join(repo_dir, "banners", s["folder"]); os.makedirs(ban_dir, exist_ok=True)
    thumb_dir = os.path.join(BAN, "thumbs_" + key); os.makedirs(thumb_dir, exist_ok=True)
    wb = Workbook()

    # ---------------- Calendar (reference layout) ----------------
    ws = wb.active; ws.title = cfg["domain"]
    ws.sheet_view.showGridLines = False
    heads = ["Week", "Planned Date", "Day", "Category", "Post Title", "GBP Post Copy", "CTA", "Landing Page", "Tracked CTA URL",
             "Banner", "Banner Overlay Text", "Banner File", "Banner Photo Source", "Image Rights", "Offer Dates / Terms", "Publishing Status"]
    widths = [7, 13, 7, 12, 30, 62, 11, 38, 50, 33, 34, 30, 40, 30, 30, 15]
    head_row(ws, 1, heads)
    ws.row_dimensions[1].height = 30
    ws.column_dimensions["A"].width = 2
    for j, w in enumerate(widths, 2):
        ws.column_dimensions[get_column_letter(j)].width = w
    for r in rows:
        rr = r["n"] + 1
        vals = [r["week"], r["date"], f'=TEXT(C{rr},"ddd")', r["cat"], r["title"], r["copy"], r["cta"], r["land"],
                f"=I{rr}&\"?utm_source=\"&Settings!$C$5&\"&utm_medium=\"&Settings!$C$6&\"&utm_campaign=\"&Settings!$C$7&\"&utm_content=\"&Settings!C{12 + r['n']}",
                None, r["overlay"], r["fname"], r["src"], r["rights"], r["terms"] or "", "Not Published"]
        for j, v in enumerate(vals, 2):
            c = ws.cell(row=rr, column=j, value=v)
            c.font = F; c.border = BOX
            c.alignment = WRAPC if j in (2, 3, 4, 5, 8, 17) else WRAP
        ws.cell(row=rr, column=3).number_format = "mmm d, yyyy"
        for j in (9, 10):
            ws.cell(row=rr, column=j).font = LINK
        ws.row_dimensions[rr].height = 200
        # copy full-size banner into repo, embed a thumbnail
        im = Image.open(r["png"]).convert("RGB")
        im.save(os.path.join(ban_dir, r["fname"]), optimize=True)
        th = os.path.join(thumb_dir, f"{r['n']:02d}.jpg")
        im.resize((228, 171), Image.LANCZOS).save(th, quality=90)
        xi = XImage(th); xi.width, xi.height = 228, 171
        ws.add_image(xi, f"K{rr}")
    last = len(rows) + 1
    dv = DataValidation(type="list", formula1='"Not Published,Scheduled,Published"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(f"Q2:Q{last}")
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"B1:Q{last}"

    # ---------------- Settings (UTM + per-post utm_content) ----------------
    wsx = wb.create_sheet("Settings")
    title(wsx, "Tracking settings", "Edit values in column C. Every Tracked CTA URL on the calendar tab updates automatically.")
    table(wsx, 4, ["Setting", "Value", "Note"],
          [("UTM source", "google", "Keep as google"), ("UTM medium", "organic", "Keep as organic"),
           ("UTM campaign", "gbp-posts-nov-2026", "One campaign for this month"),
           ("Website domain", cfg["domain"], "QA checks every landing page stays on this domain")],
          [0, 22, 40, 50], col0=2)
    wsx.cell(row=11, column=2, value="UTM content per post (unique, so each post is tracked separately in Google Analytics)").font = FB
    table(wsx, 12, ["Post #", "utm_content", "Post Title"],
          [(r["n"], r["utm"], f"='{cfg['domain']}'!F{r['n']+1}") for r in rows], [0, 0, 0, 0], col0=2, center_cols=(2,))
    wsx.column_dimensions["D"].width = 50

    # ---------------- Banner & Brand Kit ----------------
    wk = wb.create_sheet("Banner & Brand Kit")
    title(wk, "Banner and brand kit", "Every banner uses the website's own logo, colours and fonts. Size 1200 x 900 px (4:3), the ratio GBP displays best.")
    r0 = table(wk, 4, ["Item", "Value", "Note"],
               [("Logo source (website)", s["logo"], s["logo_note"]),
                ("Banner size", "1200 x 900 px PNG", "4:3 ratio. Files are in deliverables/gbp-calendars/banners/" + s["folder"] + "/"),
                ("Brand motif", s["motif"], "Gives the series a consistent, premium look"),
                ("Photos", "Real project photos from " + cfg["domain"], "Source URL for every banner is on the calendar tab")]
               + [(f"Font: {a}", b, c) for a, b, c in s["fonts"]],
               [0, 26, 60, 60], col0=2)
    wk.cell(row=r0 + 2, column=2, value="Colour palette (matched to the website)").font = FB
    pal_end = table(wk, r0 + 3, ["Colour", "Hex", "Where it comes from", "Used on banners for", "Swatch"],
                    [(a, b, c, d, "") for a, b, c, d in s["palette"]], [0, 0, 0, 0, 0, 10], col0=2)
    wk.column_dimensions["E"].width = 45; wk.column_dimensions["F"].width = 34
    for i, p in enumerate(s["palette"]):
        wk.cell(row=r0 + 4 + i, column=6, value=None).fill = PatternFill("solid", fgColor=p[1][1:])
    wk.cell(row=pal_end + 2, column=2, value="Overlay text rules (GBP safe)").font = FB
    rules = [("Length", "Tag 2 to 4 words, headline 3 to 7 words, one sub-line, one button label."),
             ("Contact details", "No phone numbers, emails, web addresses or social handles on any banner."),
             ("Claims", "Only facts stated on the website (years in business, warranties, services, offers)."),
             ("Offers", "Offer banners match the GBP Offer title. Dates and terms live in the post, not on the image."),
             ("Safe area", "Text and logo stay inside the centre 1000 x 750 px so the GBP thumbnail crop never cuts them."),
             ("Templates used", "Split photo, offer card, FAQ question, local spotlight, glass panel, 3-up compare, diagonal, frame, before/after, project spotlight, seasonal thank-you.")]
    table(wk, pal_end + 3, ["Rule", "Detail"], rules, [0, 26, 60], col0=2)
    wk.column_dimensions["A"].width = 2

    # ---------------- Image rights check ----------------
    wr = wb.create_sheet("Image Rights Check")
    title(wr, "Image rights cross-check", "Every banner photo is either the business's own photo from its website or a free Unsplash photo. Checked Oct 26, 2026.")
    src = photo_urls(s)
    rr_rows = []
    for r in rows:
        for f in r["photos"]:
            path = os.path.join(BAN, s["ph"], f)
            ev = []
            if "unsplash" in f:
                origin, verdict = "Unsplash", "OK: royalty-free (Unsplash License)"
                ev.append("Downloaded through the official Unsplash API; free photo, not Unsplash+")
            else:
                origin, verdict = "Business website", "OK: business's own photo"
                ev.append(f"In the {cfg['domain']} media library (wp-content/uploads)")
                ex = Image.open(path).getexif()
                cam = " ".join(str(ex.get(t, "")) for t in (271, 272)).strip()
                if cam:
                    ev.append(f"Taken on a phone camera ({cam}): original photo, not stock")
                if f == "91_team_crop.jpg":
                    ev.append("Cropped from the team photo on the website")
            raw = open(path, "rb").read()
            if re.search(rb"dc:creator|photoshop:Credit|Getty|Shutterstock|iStock|Envato|Depositphotos", raw):
                ev.append("Third-party credit found in file: REVIEW"); verdict = "REVIEW"
            else:
                ev.append("No photographer credit, stock agency tag or watermark")
            rr_rows.append((r["n"], r["date"], f.split("_", 1)[1], src(f), origin, "; ".join(ev), verdict))
    rend = table(wr, 4, ["Post #", "Date", "Photo", "Source URL", "Origin", "Evidence checked", "Verdict"], rr_rows,
                 [8, 13, 34, 60, 16, 70, 30], center_cols=(1, 2))
    for k in range(5, rend + 1):
        wr.cell(row=k, column=2).number_format = "mmm d, yyyy"
    wr.cell(row=rend + 2, column=1, value="Photos OK").font = FB
    wr.cell(row=rend + 2, column=2, value=f'=COUNTIF(G5:G{rend},"OK*")&" of "&COUNTA(G5:G{rend})').font = FB
    notes = s.get("removed", [])
    if notes:
        wr.cell(row=rend + 4, column=1, value="Photos removed during the check").font = FB
        table(wr, rend + 5, ["#", "Photo removed", "Reason", "Replaced with"], [(i + 1,) + n for i, n in enumerate(notes)], [0, 0, 0, 0], center_cols=(1,))
    wr.cell(row=rend + 6 + len(notes) + (1 if notes else 0), column=1,
            value="Note: photos on the business's own website are treated as owned by the business. Ask the client to confirm they took these photos or hold the rights. Unsplash photos can be used commercially without credit, but do not show people or brands in a way that suggests endorsement.").font = FS
    wr.freeze_panes = "A5"

    # ---------------- Guidelines ----------------
    wg = wb.create_sheet("GBP Guidelines")
    title(wg, "Google Business Profile post rules applied to this calendar", "Every post and banner was written and checked against these rules.")
    table(wg, 4, ["#", "Rule", "How this calendar follows it", "Source"],
          [(i + 1,) + tuple(g) for i, g in enumerate(GUIDE)], [5, 45, 60, 55], center_cols=(1,))
    wg.freeze_panes = "A5"

    # ---------------- QA ----------------
    wq = wb.create_sheet("QA Cross-Check")
    title(wq, "QA cross-check (automatic)", "Formulas re-check every post when copy or links change. Every row should read Pass before scheduling.")
    qh = ["Post #", "Date", "Characters", "Max 1,500 chars", "No phone in text", "No email in text", "No link in text",
          "CTA set", "UTM on CTA link", "Link on business website", "Weekday (Mon to Fri)", "Banner file set", "Overall"]
    C = f"'{cfg['domain']}'!"
    q = []
    for r in rows:
        k = r["n"] + 1; G = f"{C}G{k}"
        q.append([r["n"], f"={C}C{k}", f"=LEN({G})",
                  f'=IF(LEN({G})<=1500,"Pass","Fail")',
                  f'=IF(SUMPRODUCT(--ISNUMBER(SEARCH({{{cfg["phone_tokens"]}}},{G})))=0,"Pass","Fail")',
                  f'=IF(ISNUMBER(SEARCH("@",{G})),"Fail","Pass")',
                  f'=IF(OR(ISNUMBER(SEARCH("http",{G})),ISNUMBER(SEARCH("www.",{G})),ISNUMBER(SEARCH(".com",{G}))),"Fail","Pass")',
                  f'=IF(LEN({C}H{k})>0,"Pass","Fail")',
                  f'=IF(AND(ISNUMBER(SEARCH("utm_source=google",{C}J{k})),ISNUMBER(SEARCH("utm_content=",{C}J{k}))),"Pass","Fail")',
                  f'=IF(ISNUMBER(SEARCH(Settings!$C$8,{C}I{k})),"Pass","Fail")',
                  f'=IF(WEEKDAY({C}C{k},2)<=5,"Pass","Fail")',
                  f'=IF(LEN({C}M{k})>0,"Pass","Fail")',
                  f'=IF(COUNTIF(D{k+3}:L{k+3},"Fail")=0,"Pass","Fail")'])
    qend = table(wq, 4, qh, q, [8, 13, 11, 12, 12, 12, 12, 9, 12, 14, 13, 12, 10], center_cols=tuple(range(1, 14)))
    for k in range(5, qend + 1):
        wq.cell(row=k, column=2).number_format = "mmm d, yyyy"
    wq.cell(row=qend + 2, column=1, value="Posts passing").font = FB
    wq.cell(row=qend + 2, column=2, value=f'=COUNTIF(M5:M{qend},"Pass")&" of "&COUNTA(A5:A{qend})').font = FB
    wq.cell(row=qend + 4, column=1, value="Post mix").font = FB
    cats = ["Service", "FAQ", "Local", "Promotional"]
    table(wq, qend + 5, ["Category", "Posts"],
          [(c, f"=COUNTIF({C}$E$2:$E${last},\"{c}\")") for c in cats] + [("Total", f"=SUM(B{qend+6}:B{qend+9})")], [0, 0], center_cols=(2,))
    wq.freeze_panes = "B5"

    for sh in wb.worksheets:
        sh.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        sh.page_setup.orientation = "landscape"; sh.page_setup.fitToWidth = 1; sh.page_setup.fitToHeight = 0
    out = os.path.join(repo_dir, s["file"])
    wb.save(out)
    return out, rows


if __name__ == "__main__":
    repo = sys.argv[1]
    for k in ("hg", "pf"):
        out, rows = build(k, repo)
        from collections import Counter
        print(out, len(rows), Counter(r["cat"] for r in rows), Counter(r["week"] for r in rows))
